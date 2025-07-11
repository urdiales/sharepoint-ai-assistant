"""
Helper functions for parsing, file preview, etc.
This module provides utilities for previewing different file types from SharePoint.
"""

import pandas as pd
import io
from pdfplumber import open as open_pdf
from docx import Document
import openpyxl

def preview_pdf(file_bytes):
    # Preview PDF file contents (first page text)
    # Takes a BytesIO object containing PDF data and extracts text from the first page
    with open_pdf(file_bytes) as pdf:
        page = pdf.pages[0]  # Get the first page of the PDF
        return page.extract_text()  # Extract and return the text content

def preview_docx(file_bytes):
    # Preview DOCX file contents (first 1000 chars)
    # Takes a BytesIO object containing DOCX data and extracts text from paragraphs
    doc = Document(file_bytes)
    full_text = []
    
    # Extract text from each paragraph in the document
    for para in doc.paragraphs:
        full_text.append(para.text)
        
    # Join paragraphs with newlines and limit to first 1000 characters
    return "\n".join(full_text)[:1000]

def preview_xlsx(file_bytes):
    # Preview XLSX file as DataFrame (first 10 rows)
    # Takes a BytesIO object containing XLSX data and converts to a pandas DataFrame
    wb = openpyxl.load_workbook(file_bytes, read_only=True)
    ws = wb.active  # Get the active worksheet
    
    # Extract data from the first 10 rows
    data = [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=10)]
    
    # Create DataFrame using first row as column headers and remaining rows as data
    df = pd.DataFrame(data[1:], columns=data[0])
    return df
